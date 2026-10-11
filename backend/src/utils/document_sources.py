"""Location-aware extraction of original lecture sources."""
import base64
from io import BytesIO
import re

import pymupdf
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.utils.source_contracts import FILE_TYPES, MAX_SOURCE_CHARACTERS, MAX_SOURCE_UNITS

OFFICE_WARNING = 'Unsupported Office visual content; convert the original to PDF to preserve visual fidelity.'
R_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


def _file_type(file_type):
    if not isinstance(file_type, str) or file_type.lower().lstrip('.') not in FILE_TYPES:
        raise ValueError('Unsupported source file type')
    return file_type.lower().lstrip('.')


def _unit(file_type, position, text='', assets=None, warnings=None, label=None):
    prefix, kind = FILE_TYPES[file_type]
    assets, warnings = assets or [], list(dict.fromkeys(warnings or []))
    visual = bool(assets or warnings)
    return {'unitId': f'{prefix}{position:04d}', 'kind': kind, 'position': position,
            'label': label or f'{kind.title()} {position}', 'text': text,
            'visualAssets': assets, 'needsVisualRead': visual,
            'visualReadStatus': 'unsupported' if warnings else ('pending' if visual else 'not_required'),
            'blank': not bool(text.strip() or visual), 'warnings': warnings}


def _limits(units, file_type):
    if file_type != 'docx' and len(units) > MAX_SOURCE_UNITS:
        raise ValueError('Split the file: maximum 160 pages/slides per source')
    if sum(len(u['text']) for u in units) > MAX_SOURCE_CHARACTERS:
        raise ValueError('Split the file: maximum 600,000 extracted characters per source')
    if not units:
        raise ValueError('Original source contains no addressable units')
    return units


def _pdf_units(raw):
    units = []
    with pymupdf.open(stream=raw, filetype='pdf') as doc:
        if doc.needs_pass:
            raise ValueError('Password-protected original cannot be read')
        if len(doc) > MAX_SOURCE_UNITS:
            raise ValueError('Split the file: maximum 160 pages/slides per source')
        for position, page in enumerate(doc, 1):
            visual = bool(page.get_image_info() or page.get_drawings() or list(page.annots() or []))
            # A full original page is the addressable asset, including vector art.
            assets = [{'assetId': f'p{position:04d}-page', 'kind': 'page', 'mimeType': 'image/png',
                       'unitId': f'p{position:04d}'}] if visual else []
            units.append(_unit('pdf', position, page.get_text('text', sort=True), assets))
    return units


def _office_assets(element, part, asset_prefix):
    """Read only embedded relationships; never dereference external resources."""
    assets, warnings = [], []
    for node in element.iter():
        local = node.tag.rsplit('}', 1)[-1]
        if node.tag.startswith('{urn:schemas-microsoft-com:vml}') and local != 'imagedata':
            warnings.append(OFFICE_WARNING)
        if local == 'blip' or local == 'imagedata':
            rid = node.get(f'{{{R_NS}}}embed') or node.get(f'{{{R_NS}}}id')
            linked = node.get(f'{{{R_NS}}}link')
            rel = part.rels.get(rid) if rid else None
            if linked or not rel or rel.is_external or '/image' not in rel.reltype:
                warnings.append(OFFICE_WARNING)
                continue
            blob = rel.target_part.blob
            try:
                with Image.open(BytesIO(blob)) as img:
                    img.load()
            except Exception:
                warnings.append(OFFICE_WARNING)
                continue
            assets.append({'assetId': f'{asset_prefix}-image{len(assets)+1:04d}', 'kind': 'image',
                           'mimeType': rel.target_part.content_type, 'relationshipId': rid,
                           'partName': str(part.partname)})
        elif local in ('chart', 'relIds', 'oleObj', 'AlternateContent', 'shape', 'object', 'oMath', 'oMathPara'):
            warnings.append(OFFICE_WARNING)
        elif local == 'graphicData' and not any(kind in node.get('uri', '') for kind in ('picture', '/table')):
            warnings.append(OFFICE_WARNING)
    return assets, warnings


def _slide_content(shapes, inherited=False):
    texts, warnings = [], []
    for shape in shapes:
        if inherited and shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            nested, nested_warnings = _slide_content(shape.shapes, inherited=inherited)
            texts.extend(nested)
            warnings.extend(nested_warnings)
        elif shape.has_table:
            texts.append('\n'.join('\t'.join(cell.text for cell in row.cells) for row in shape.table.rows))
        elif shape.has_text_frame:
            texts.append(shape.text_frame.text)
            # Auto-shapes, connectors and freeforms carry geometry we cannot render.
            if shape.shape_type not in (MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PLACEHOLDER):
                warnings.append(OFFICE_WARNING)
        elif shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            warnings.append(OFFICE_WARNING)
    return texts, warnings


def _pptx_units(raw):
    prs = Presentation(BytesIO(raw))
    if len(prs.slides) > MAX_SOURCE_UNITS:
        raise ValueError('Split the file: maximum 160 pages/slides per source')
    units = []
    for position, slide in enumerate(prs.slides, 1):
        texts, warnings = _slide_content(slide.shapes)
        assets, asset_warnings = _office_assets(slide.element, slide.part, f's{position:04d}')
        warnings.extend(asset_warnings)
        # Layout/master backgrounds and fixed artwork appear in the original
        # slide even when there is no corresponding slide-local picture shape.
        for story_name, story in [('layout', slide.slide_layout), ('master', slide.slide_layout.slide_master)]:
            extra, inherited_warnings = _office_assets(story.element, story.part, f's{position:04d}-{story_name}')
            inherited_texts, shape_warnings = _slide_content(story.shapes, inherited=True)
            assets.extend(extra)
            texts.extend(inherited_texts)
            warnings.extend(inherited_warnings + shape_warnings)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text:
                texts.append(frame.text)
            extra, note_warnings = _office_assets(slide.notes_slide.element, slide.notes_slide.part, f's{position:04d}-notes')
            assets.extend(extra)
            warnings.extend(note_warnings)
        units.append(_unit('pptx', position, '\n'.join(texts), assets, warnings))
    return units


def _block_text(element):
    # Paragraph boundaries and cell boundaries are retained, including nested tables.
    if element.tag == qn('w:tbl'):
        return '\n'.join('\t'.join('\n'.join(_block_text(child) for child in cell
                                            if child.tag in (qn('w:p'), qn('w:tbl')))
                                   for cell in row if cell.tag == qn('w:tc'))
                         for row in element if row.tag == qn('w:tr'))
    if element.tag == qn('w:p'):
        pieces = []
        for node in element.iter():
            if node.tag in (qn('w:t'), qn('m:t')):
                pieces.append(node.text or '')
            elif node.tag == qn('w:tab'):
                pieces.append('\t')
            elif node.tag in (qn('w:br'), qn('w:cr')):
                pieces.append('\n')
        return ''.join(pieces)
    return '\n'.join(_block_text(node) for node in element.iter(qn('w:p')))


def _docx_blocks(container):
    for element in container:
        if element.tag in (qn('w:p'), qn('w:tbl')):
            yield element
        elif element.tag != qn('w:sectPr'):
            # Content controls and tracked insertions still contain source blocks.
            children = list(_docx_blocks(element))
            if children:
                yield from children
            else:
                yield element


def _docx_units(raw):
    doc = Document(BytesIO(raw))
    stories = [(doc.element.body, doc.part, 'Body')]
    seen = set()
    for section_no, section in enumerate(doc.sections, 1):
        for name in ('header', 'first_page_header', 'even_page_header', 'footer', 'first_page_footer', 'even_page_footer'):
            story = getattr(section, name)
            if story.is_linked_to_previous:
                continue
            if str(story.part.partname) in seen:
                continue
            seen.add(str(story.part.partname))
            stories.append((story._element, story.part, f'Section {section_no} {name.replace("_", " ")}'))
    units = []
    for container, part, story_name in stories:
        for block_no, element in enumerate(_docx_blocks(container), 1):
            position = len(units) + 1
            assets, warnings = _office_assets(element, part, f'b{position:04d}')
            if element.tag not in (qn('w:p'), qn('w:tbl')):
                warnings.append('Unsupported DOCX block; convert the original to PDF to preserve content.')
            # Textboxes have their own paragraphs, not physical page coordinates.
            text = _block_text(element)
            block_kind = 'table' if element.tag == qn('w:tbl') else 'paragraph'
            units.append(_unit('docx', position, text, assets, warnings,
                               f'{story_name} {block_kind} {block_no}'))
    return units


def extract_source_units(file_bytes: bytes, file_type: str) -> list[dict]:
    file_type = _file_type(file_type)
    if not isinstance(file_bytes, bytes) or not file_bytes:
        raise ValueError('Original source bytes are empty or invalid')
    try:
        units = {'pdf': _pdf_units, 'pptx': _pptx_units, 'docx': _docx_units}[file_type](file_bytes)
        return _limits(units, file_type)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f'Cannot read the original {file_type.upper()} source') from exc


def _image_uri(blob):
    with Image.open(BytesIO(blob)) as img:
        img.load()
        out = BytesIO()
        img.convert('RGB').save(out, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(out.getvalue()).decode('ascii')


def render_source_unit(file_bytes: bytes, file_type: str, unit_id: str) -> list[str]:
    file_type = _file_type(file_type)
    prefix = FILE_TYPES[file_type][0]
    if not isinstance(unit_id, str) or not re.fullmatch(prefix + r'\d{4,}', unit_id):
        raise ValueError('Invalid source location')
    position = int(unit_id[1:])
    if position < 1 or unit_id != f'{prefix}{position:04d}':
        raise ValueError('Invalid source location')
    if file_type == 'pdf':
        try:
            with pymupdf.open(stream=file_bytes, filetype='pdf') as doc:
                if doc.needs_pass or position > len(doc):
                    raise ValueError('Source page unavailable')
                pix = doc[position-1].get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
                return ['data:image/png;base64,' + base64.b64encode(pix.tobytes('png')).decode('ascii')]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError('Cannot render original PDF page') from exc
    units = extract_source_units(file_bytes, file_type)
    if position > len(units):
        raise ValueError('Source location unavailable')
    original = Presentation(BytesIO(file_bytes)) if file_type == 'pptx' else Document(BytesIO(file_bytes))
    parts = {str(part.partname): part for part in original.part.package.iter_parts()}
    result = []
    for asset in units[position-1]['visualAssets']:
        rel = parts[asset['partName']].rels[asset['relationshipId']]
        if rel.is_external:
            raise ValueError('External source images are unsupported')
        result.append(_image_uri(rel.target_part.blob))
    return result


def chunk_source_unit(unit: dict, max_characters: int = 6000, overlap: int = 400) -> list[dict]:
    if type(max_characters) is not int or type(overlap) is not int or not 0 <= overlap < max_characters:
        raise ValueError('Invalid source chunk size or overlap')
    if not isinstance(unit, dict) or not isinstance(unit.get('text'), str) or not isinstance(unit.get('unitId'), str):
        raise ValueError('Invalid source chunk unit')
    text, chunks, start = unit['text'], [], 0
    while start < len(text):
        end = min(len(text), start + max_characters)
        chunks.append({'chunkId': f'{unit["unitId"]}-c{len(chunks)+1:04d}', 'unitId': unit['unitId'],
                       'start': start, 'end': end, 'text': text[start:end]})
        if end == len(text):
            break
        start = end - overlap
    return chunks
