"""Synthetic originals catch lost late units, text, and Office content."""
import base64
from io import BytesIO
from lxml import etree

import fitz
from PIL import Image
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import pytest

from src.utils import document_sources as sources


def png_bytes():
    out = BytesIO()
    Image.new('RGB', (30, 20), 'red').save(out, 'PNG')
    return out.getvalue()


def pdf_bytes(count=17):
    doc = fitz.open()
    for i in range(count):
        page = doc.new_page()
        if i != 2:
            page.insert_text((40, 40), 'LATE_FORMULA_SENTINEL' if i == count - 1 else f'PAGE_{i+1}')
        if i in (0, 1):
            page.insert_image(fitz.Rect(40, 70, 100, 110), stream=png_bytes())
        if i == 3:
            # Actual visible vector content must be read even without raster images.
            page.draw_rect(fitz.Rect(40, 70, 100, 110), color=(1, 0, 0))
    result = doc.tobytes()
    doc.close()
    return result


def pptx_bytes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = 'EARLY_SLIDE_TEXT'
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(3), Inches(1)).table
    table.cell(0, 0).text = 'TABLE_SENTINEL'
    table.cell(0, 1).text = 'SECOND_CELL'
    slide.notes_slide.notes_text_frame.text = 'NOTES_SENTINEL'
    slide.shapes.add_picture(BytesIO(png_bytes()), Inches(1), Inches(3))
    prs.slides.add_slide(prs.slide_layouts[6])
    image_slide = prs.slides.add_slide(prs.slide_layouts[6])
    image_slide.shapes.add_picture(BytesIO(png_bytes()), Inches(1), Inches(1))
    chart_slide = prs.slides.add_slide(prs.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['A', 'B']
    data.add_series('S', (1, 2))
    chart_slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(3), Inches(2), data)
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def docx_bytes():
    doc = Document()
    doc.add_heading('HEADING_SENTINEL', 1)
    doc.add_paragraph('BEFORE_TABLE')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'DOC_TABLE'
    table.cell(0, 1).text = 'CELL_TWO'
    doc.add_paragraph('AFTER_TABLE')
    doc.add_picture(BytesIO(png_bytes()))
    doc.sections[0].header.paragraphs[0].text = 'HEADER_SENTINEL'
    doc.sections[0].footer.paragraphs[0].text = 'FOOTER_SENTINEL'
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def test_pdf_keeps_all_late_blank_and_visual_pages():
    units = sources.extract_source_units(pdf_bytes(), 'pdf')
    assert [u['unitId'] for u in units] == [f'p{i:04d}' for i in range(1, 18)]
    assert 'LATE_FORMULA_SENTINEL' in units[-1]['text']
    assert units[2]['blank'] is True
    assert units[0]['needsVisualRead'] is True
    assert units[3]['needsVisualRead'] is True
    assert units[3]['blank'] is False


def test_pdf_image_only_page_and_exact_page_render():
    doc = fitz.open()
    doc.new_page().insert_image(fitz.Rect(10, 10, 40, 30), stream=png_bytes())
    raw = doc.tobytes()
    doc.close()
    unit = sources.extract_source_units(raw, 'pdf')[0]
    assert unit['text'] == '' and unit['needsVisualRead'] and not unit['blank']
    rendered = sources.render_source_unit(pdf_bytes(), 'pdf', 'p0017')
    image = Image.open(BytesIO(base64.b64decode(rendered[0].split(',', 1)[1])))
    assert image.width > 500
    with pytest.raises(ValueError):
        sources.render_source_unit(pdf_bytes(), 'pdf', 'p0018')
    with pytest.raises(ValueError):
        sources.render_source_unit(pdf_bytes(), 'pdf', '../p0001')


def test_pptx_keeps_tables_notes_images_blank_and_unresolved_charts():
    raw = pptx_bytes()
    units = sources.extract_source_units(raw, 'pptx')
    assert [u['unitId'] for u in units] == ['s0001', 's0002', 's0003', 's0004']
    assert all(value in units[0]['text'] for value in ['EARLY_SLIDE_TEXT', 'TABLE_SENTINEL', 'SECOND_CELL', 'NOTES_SENTINEL'])
    assert units[1]['blank'] is True
    assert units[2]['needsVisualRead'] and not units[2]['blank']
    assert units[3]['warnings'] and not units[3]['blank']
    assert 'PDF' in ' '.join(units[3]['warnings'])
    assert sources.render_source_unit(raw, 'pptx', 's0003')[0].startswith('data:image/')
    assert sources.render_source_unit(raw, 'pptx', 's0002') == []


def test_docx_keeps_order_headers_footers_and_embedded_images_without_pages():
    raw = docx_bytes()
    units = sources.extract_source_units(raw, 'docx')
    texts = [u['text'] for u in units]
    assert texts.index('BEFORE_TABLE') < next(i for i, t in enumerate(texts) if 'DOC_TABLE' in t) < texts.index('AFTER_TABLE')
    assert any('HEADER_SENTINEL' in t for t in texts)
    assert any('FOOTER_SENTINEL' in t for t in texts)
    assert 'DOC_TABLE\tCELL_TWO' in texts
    assert all(u['kind'] == 'block' and 'page' not in u['label'].lower() for u in units)
    pictured = next(u for u in units if u['visualAssets'])
    assert sources.render_source_unit(raw, 'docx', pictured['unitId'])[0].startswith('data:image/')


def test_office_math_is_preserved_as_text_with_unresolved_fidelity_warning():
    doc = Document()
    paragraph = doc.add_paragraph('FORMULA: ')
    math = OxmlElement('m:oMath')
    run = OxmlElement('m:r')
    text = OxmlElement('m:t')
    text.text = 'x=3'
    run.append(text)
    math.append(run)
    paragraph._p.append(math)
    out = BytesIO()
    doc.save(out)
    unit = sources.extract_source_units(out.getvalue(), 'docx')[0]
    assert 'x=3' in unit['text']
    assert unit['warnings'] and unit['visualReadStatus'] == 'unsupported'


def test_docx_legacy_vector_drawing_is_not_misclassified_as_a_blank_paragraph():
    doc = Document()
    paragraph = doc.add_paragraph()
    pict = OxmlElement('w:pict')
    pict.append(etree.Element('{urn:schemas-microsoft-com:vml}rect'))
    paragraph._p.append(pict)
    out = BytesIO()
    doc.save(out)
    unit = sources.extract_source_units(out.getvalue(), 'docx')[0]
    assert not unit['blank']
    assert unit['visualReadStatus'] == 'unsupported'


def test_chunks_cover_every_character_with_stable_ranges_and_overlap():
    unit = {'unitId': 'b0001', 'text': ('عربي😀\nformula x=3; ' * 1300)}
    chunks = sources.chunk_source_unit(unit, max_characters=600, overlap=40)
    assert chunks[0]['start'] == 0 and chunks[-1]['end'] == len(unit['text'])
    assert all(c['text'] == unit['text'][c['start']:c['end']] and len(c['text']) <= 600 for c in chunks)
    assert all(b['start'] == a['end'] - 40 for a, b in zip(chunks, chunks[1:]))
    assert sum(len(c['text']) for c in chunks) >= len(unit['text'])
    with pytest.raises(ValueError):
        sources.chunk_source_unit(unit, max_characters=40, overlap=40)


@pytest.mark.parametrize('file_type', ['pdf', 'pptx', 'docx', 'exe'])
def test_corrupt_or_unsupported_files_fail_explicitly(file_type):
    with pytest.raises(ValueError):
        sources.extract_source_units(b'not a valid original', file_type)


def test_pdf_limit_fails_without_returning_a_truncated_source():
    assert len(sources.extract_source_units(pdf_bytes(160), 'pdf')) == 160
    with pytest.raises(ValueError, match='[Ss]plit'):
        sources.extract_source_units(pdf_bytes(161), 'pdf')


def test_text_limit_counts_complete_docx_source():
    doc = Document()
    doc.add_paragraph('x' * 600001)
    out = BytesIO()
    doc.save(out)
    with pytest.raises(ValueError, match='[Ss]plit'):
        sources.extract_source_units(out.getvalue(), 'docx')


def test_docx_external_image_and_vector_are_unresolved_without_network_access():
    doc = Document()
    paragraph = doc.add_paragraph('VISIBLE_TEXT')
    drawing = OxmlElement('w:drawing')
    blip = OxmlElement('a:blip')
    blip.set(qn('r:link'), 'rIdExternalImage')
    drawing.append(blip)
    paragraph._p.append(drawing)
    out = BytesIO()
    doc.save(out)
    unit = sources.extract_source_units(out.getvalue(), 'docx')[0]
    assert unit['text'] == 'VISIBLE_TEXT'
    assert unit['visualReadStatus'] == 'unsupported'
    assert unit['visualAssets'] == []
    assert unit['warnings'] and unit['needsVisualRead']


def test_docx_more_than_160_blocks_is_not_misreported_as_physical_pages():
    doc = Document()
    for i in range(170):
        doc.add_paragraph(f'CONTENT_{i}')
    out = BytesIO()
    doc.save(out)
    assert len(sources.extract_source_units(out.getvalue(), 'docx')) == 170


def test_pptx_supported_table_does_not_require_an_unsupported_visual_resolution():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1)).table.cell(0, 0).text = 'ONLY_TABLE'
    out = BytesIO()
    prs.save(out)
    unit = sources.extract_source_units(out.getvalue(), 'pptx')[0]
    assert unit['text'] == 'ONLY_TABLE'
    assert unit['warnings'] == []
    assert unit['needsVisualRead'] is False


def test_pptx_inherited_layout_image_is_preserved_and_renderable():
    prs = Presentation()
    layout = prs.slide_layouts[6]
    # python-pptx does not expose add_picture for layouts; move a real picture
    # into the layout and transfer its embedded relationship.
    slide = prs.slides.add_slide(layout)
    picture = slide.shapes.add_picture(BytesIO(png_bytes()), Inches(1), Inches(1))
    image_part = slide.part.rels[picture._pic.blipFill.blip.rEmbed].target_part
    rid = layout.part.relate_to(image_part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image')
    picture._pic.blipFill.blip.rEmbed = rid
    layout.shapes._spTree.append(picture._pic)
    out = BytesIO()
    prs.save(out)
    raw = out.getvalue()
    unit = sources.extract_source_units(raw, 'pptx')[0]
    assert unit['visualAssets'] and not unit['blank']
    assert sources.render_source_unit(raw, 'pptx', 's0001')[0].startswith('data:image/png;base64,')
