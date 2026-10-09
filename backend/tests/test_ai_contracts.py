import json
from pathlib import Path
import sys
from types import SimpleNamespace
from unittest.mock import patch

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.utils.json_parse import extract_json_object, parse_with_retry
from src.utils.topic_scope import course_scope_from_docs, topic_in_scope, retag_topic_weights
from src.utils.ai_retry import chat_with_retry
from src.agents.tutor_agent import AITutorAgent
from src.agents.exam_generator import ExamGeneratorAgent
from src.utils.pdf_extract import image_to_image_uri
from src.utils.ai_contracts import validate_analysis, validate_summary, validate_exam_contract
from src.agents.document_processor import DocumentProcessorAgent


def test_truncated_json_rejects_finished_prefix_instead_of_silent_loss():
    with pytest.raises(ValueError):
        extract_json_object('{"problems":[{"statement":"first"},{"statement":"cut')


def test_missing_numeric_field_retries_then_fails_instead_of_defaulting_list():
    calls = []
    def call(reminder):
        calls.append(reminder)
        return '{"topics": ["trees"], "problems": []}'
    with pytest.raises(RuntimeError):
        parse_with_retry(call, required_keys=('topics', 'problems', 'problemCount'))
    assert len(calls) == 2


def test_parser_errors_do_not_include_private_model_content():
    with pytest.raises(ValueError) as error:
        extract_json_object('private-course-note-token')
    assert 'private-course-note-token' not in str(error.value)


def test_unicode_scope_supports_arabic_and_legacy_dictionary_topics():
    _, words = course_scope_from_docs([{'topics': [{'title': 'الشبكات العصبية'}, {'name': 'Decision Trees'}]}])
    assert topic_in_scope('الشبكات العصبية', words)
    assert not topic_in_scope('قواعد البيانات', words)
    assert topic_in_scope({'topic': 'Decision tree'}, words)
    assert not topic_in_scope({'unexpected': 'private'}, words)


def test_unrecognizable_document_scope_cannot_admit_every_historical_topic():
    tagged = retag_topic_weights([{'topicWeights': [{'topic': 'Databases', 'weight': 1}]}], [{'topics': ['!!!']}])
    assert tagged[0]['topicWeights'][0]['inScope'] is False


def test_historical_only_generation_clears_stale_document_scope_flags():
    historical = [{'topicWeights': [{'topic': 'Trees', 'inScope': False}], 'scopeChecked': True, 'outOfCourseTopics': ['Trees']}]
    tagged = retag_topic_weights(historical, [])
    assert tagged[0]['topicWeights'][0]['inScope'] is None
    assert tagged[0]['scopeStatus'] == 'no_doc_scope'
    assert tagged[0]['outOfCourseTopics'] == []
    assert historical[0]['topicWeights'][0]['inScope'] is False


def test_tutor_context_covers_every_selected_resource_with_legacy_shapes():
    agent = AITutorAgent.__new__(AITutorAgent)
    resources = {'documents': [{'title': f'Lecture {i}', 'analysis': {'topics': [{'title': f'SENTINEL{i}'}], 'definitions': [{'term': 'x'}]},
                               'extractedText': f'UNIQUE{i} ' * 3000} for i in range(9)]}
    context = agent._build_context(resources)
    for i in range(9):
        assert f'UNIQUE{i}' in context
    assert len(context) <= 20000


def test_retry_status_classification_never_retries_permanent_private_message():
    class BadRequest(Exception): status_code = 400
    attempts = []
    def create(**kwargs):
        attempts.append(kwargs)
        raise BadRequest('private rate in exam response')
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create)))
    with patch('src.utils.ai_retry.time.sleep'):
        with pytest.raises(BadRequest): chat_with_retry(client)
    assert len(attempts) == 1


def test_provider_call_has_explicit_timeout_and_nonempty_success():
    seen = []
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=lambda **kw: seen.append(kw) or 'ok')))
    assert chat_with_retry(client) == 'ok'
    assert 0 < seen[0].get('timeout', 0) <= 120


def test_invalid_image_is_rejected_instead_of_mislabeled_jpeg():
    with pytest.raises(ValueError): image_to_image_uri(b'not an image')


def test_wrong_flashcard_count_is_rejected_instead_of_claiming_success():
    agent = ExamGeneratorAgent.__new__(ExamGeneratorAgent)
    agent.client, agent.model_id = object(), 'fake'
    response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps({'cards': [
        {'front': 'one', 'back': 'answer', 'topic': 't', 'examLikelihood': 'high'}]})), finish_reason='stop')])
    with patch('src.agents.exam_generator.chat_with_retry', return_value=response):
        with pytest.raises((RuntimeError, ValueError)):
            agent.generate_flashcards({}, [], count=3)


def test_malformed_summary_does_not_return_empty_success():
    agent = ExamGeneratorAgent.__new__(ExamGeneratorAgent)
    agent.client, agent.model_id = object(), 'fake'
    response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='private malformed output'), finish_reason='stop')])
    with patch('src.agents.exam_generator.chat_with_retry', return_value=response):
        with pytest.raises((RuntimeError, ValueError)):
            agent.generate_summary({}, [])


def test_continuation_cap_rejects_incomplete_exam_instead_of_auto_closing():
    agent = ExamGeneratorAgent.__new__(ExamGeneratorAgent)
    agent.client, agent.model_id = object(), 'fake'
    response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='\\documentclass{article}\n\\begin{document}\npartial'), finish_reason='length')])
    with patch('src.agents.exam_generator.chat_with_retry', return_value=response):
        with pytest.raises((RuntimeError, ValueError)):
            agent.compile_enhanced_exam({}, [], '', 'e')


def test_duplicate_json_question_ids_are_rejected():
    with pytest.raises(ValueError):
        extract_json_object('{"questions":{"q1":{"max_score":1},"q1":{"max_score":2}}}')


def test_trailing_comma_repair_does_not_modify_string_content():
    assert extract_json_object('{"topics":["literal ,}",],}') == {'topics': ['literal ,}']}


def test_fenced_json_preserves_code_fences_within_content():
    assert extract_json_object('```json\n{"code":"```python example ```"}\n```') == {'code': '```python example ```'}


def test_docx_tables_are_included_in_extracted_lecture_text():
    import io
    from docx import Document
    document = Document()
    document.add_paragraph('Lecture heading')
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Table-only formula E = mc2'
    buffer = io.BytesIO()
    document.save(buffer)
    agent = DocumentProcessorAgent.__new__(DocumentProcessorAgent)
    assert 'Table-only formula E = mc2' in agent.extract_text(buffer.getvalue(), 'docx')


def test_pptx_table_and_speaker_notes_are_included():
    import io
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(3), Inches(1)).table
    table.cell(0, 0).text = 'TABLESENTINEL'
    slide.notes_slide.notes_text_frame.text = 'NOTESENTINEL'
    buffer = io.BytesIO()
    presentation.save(buffer)
    agent = DocumentProcessorAgent.__new__(DocumentProcessorAgent)
    text = agent.extract_text(buffer.getvalue(), 'pptx')
    assert 'TABLESENTINEL' in text and 'NOTESENTINEL' in text


def test_transient_transcription_retry_rewinds_the_uploaded_stream():
    import io
    from src.utils.ai_retry import call_with_retry
    class Transient(Exception): status_code = 503
    attempts = []
    def transcribe(**kwargs):
        content = kwargs['file'].read()
        attempts.append(content)
        if len(attempts) == 1: raise Transient('temporary')
        return content
    with patch('src.utils.ai_retry.time.sleep'):
        result = call_with_retry(transcribe, file=io.BytesIO(b'complete audio'))
    assert result == b'complete audio'
    assert attempts == [b'complete audio', b'complete audio']


@pytest.mark.parametrize('field,value', [('keyConceptCount', []), ('keyConceptCount', True), ('topics', [{'unknown': 'x'}]), ('definitions', [{'term': 'x'}])])
def test_document_schema_rejects_invalid_nested_and_numeric_fields(field, value):
    valid = {'topics': ['trees'], 'definitions': [], 'formulas': [], 'chapterMapping': [], 'keyConceptCount': 1}
    with pytest.raises(ValueError): validate_analysis('document', {**valid, field: value})


def test_exam_contract_rejects_missing_question_types_and_printed_mark_drift():
    tex = r'\begin{document}\noindent\textbf{Question 1:} One [2 marks]\noindent\textbf{Question 2:} Two [3 marks]\end{document}'
    rubrics = {'questions': {'q1': {'question_type': 'mcq', 'max_score': 2, 'correct_answer': 'A'},
                            'q2': {'question_type': 'written', 'max_score': 3, 'criteria': 'Explain'}}}
    historical = [{'questionTypes': [{'type': 'mcq', 'count': 1}, {'type': 'written', 'count': 1}], 'totalQuestions': 2}]
    assert validate_exam_contract(tex, rubrics, historical, total_marks=5) == rubrics
    with pytest.raises(ValueError): validate_exam_contract(tex.replace('[3 marks]', '[4 marks]'), rubrics, historical, 5)
    with pytest.raises(ValueError): validate_exam_contract(tex, rubrics, [{'questionTypes': [{'type': 'mcq', 'count': 2}]}])


def test_normalized_marks_update_printed_values_without_altering_question_text():
    from src.utils import ai_contracts
    assert hasattr(ai_contracts, 'align_printed_marks')
    tex = r'\begin{document}\textbf{Question 1:} Compute $x=5$. [5 marks]\textbf{Question 2:} Prove $y=2$. [3 marks]\end{document}'
    rubrics = {'questions': {'q1': {'question_type': 'written', 'max_score': 2.5, 'criteria': 'x'}, 'q2': {'question_type': 'written', 'max_score': 1.5, 'criteria': 'y'}}}
    aligned = ai_contracts.align_printed_marks(tex, rubrics)
    assert '$x=5$' in aligned and '$y=2$' in aligned
    assert '[2.5 marks]' in aligned and '[1.5 marks]' in aligned
    assert ai_contracts.validate_exam_contract(aligned, rubrics, total_marks=4) == rubrics


def test_summary_generation_includes_every_selected_raw_document():
    agent = ExamGeneratorAgent.__new__(ExamGeneratorAgent)
    agent.client, agent.model_id = object(), 'fake'
    valid = {'title': 'T', 'overview': 'Overview', 'sections': [{'heading': 'H', 'content': 'C', 'keyPoints': [], 'examLikelihood': 'medium', 'examWeight': 100}], 'keyTerms': [], 'examFocus': []}
    response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(valid)), finish_reason='stop')])
    prompts = []
    def fake_call(*args, **kwargs):
        prompts.append(kwargs['messages'][1]['content'])
        return response
    with patch('src.agents.exam_generator.chat_with_retry', side_effect=fake_call):
        result = agent.generate_summary({}, [], document_texts=[f'SOURCE{i} ' * 3000 for i in range(9)])
    assert len(result['sections']) == 1
    for i in range(9): assert f'SOURCE{i}' in prompts[0]
    assert 'excerpted=true' in prompts[0]


def test_answer_key_rejects_naturally_finished_prefix_missing_question():
    agent = ExamGeneratorAgent.__new__(ExamGeneratorAgent)
    agent.client, agent.model_id = object(), 'fake'
    tex = r'\begin{document}\textbf{Question 1:} A\textbf{Question 2:} B\end{document}'
    response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=r'\textbf{Question 1.} A {\color{red}\textbf{Answer:} B}'), finish_reason='stop')])
    with patch('src.agents.exam_generator.chat_with_retry', return_value=response):
        with pytest.raises(ValueError): agent.generate_answer_key(tex, 'e')


def test_legacy_exam_rubric_rescue_accepts_complete_schema_without_new_heading_format():
    agent = ExamGeneratorAgent.__new__(ExamGeneratorAgent)
    agent.client, agent.model_id = object(), 'fake'
    rubrics = {'questions': {'q1': {'question_type': 'mcq', 'max_score': 2, 'correct_answer': 'A'}}}
    response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(rubrics)), finish_reason='stop')])
    tex = r'\begin{document}\begin{enumerate}\item Legacy question\end{enumerate}\end{document}'
    with patch('src.agents.exam_generator.chat_with_retry', return_value=response):
        assert agent.generate_rubrics_from_tex(tex, 'legacy') == rubrics
