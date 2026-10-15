"""Original-byte fixtures; provider calls are replaced only at the transport seam."""
import copy
import importlib
import io
import json
from types import SimpleNamespace

import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from src.supporting.contracts import validate_supporting_manifest
from src.utils.document_sources import render_source_unit
from src.utils.pdf_extract import image_to_image_uri


IDENTITY = {'userId': 'u', 'courseId': 'c', 'kind': 'historical_exam', 'materialId': 'm'}


def extractor(client=None):
    assert importlib.util.find_spec('src.supporting.extraction') is not None, 'Complete supporting extraction engine is missing'
    return importlib.import_module('src.supporting.extraction').SupportingExtractor(client, 'fixture-model')


def png(color='white', text=''):
    out = io.BytesIO()
    image = Image.new('RGB', (700, 250), color)
    ImageDraw.Draw(image).multiline_text((20, 20), text, fill='black', spacing=8)
    image.save(out, format='PNG')
    return out.getvalue()


def pdf(texts, *, scanned=False, vector=False):
    doc = fitz.open()
    for index, text in enumerate(texts):
        page = doc.new_page()
        if scanned:
            page.insert_image(fitz.Rect(30, 30, 550, 220), stream=png('white', text))
        else:
            page.insert_text((40, 50), text)
        if vector:
            page.draw_rect(fitz.Rect(50, 100, 150, 200))
    raw = doc.tobytes()
    doc.close()
    return raw


def reader():
    from src.services.source_evidence import SourceEvidence
    from .test_source_context import source_scope
    return SourceEvidence(source_scope(['Algorithms: Trees']), object(), 'u', 'c', None, 'no-lecture-model-calls')


def proof(manifest, number=1, *, visual=False, start=0, end=None):
    unit = manifest['units'][number-1]
    ref = {k: manifest['identity'][k] for k in ('kind', 'materialId')}
    ref.update(sourceRevision=manifest['sourceRevision'], unitId=unit['unitId'])
    if visual:
        return {'type': 'visual_region', 'ref': ref, 'assetId': unit['visualAssets'][0]['assetId'], 'box': [0, 0, 1, 1]}
    end = len(unit['text']) if end is None else end
    return {'type': 'text_span', 'ref': ref, 'field': 'text', 'start': start, 'end': end,
            'quote': unit['text'][start:end]}


def item(anchor, text='Explain trees', *, item_type='question', parent=None, marks=None):
    requirement = {'origin': anchor, 'text': text, 'category': 'concept', 'proofs': [anchor]}
    return {'origin': anchor, 'itemType': item_type, 'parent': parent, 'text': text,
            'printedLabel': '1' if item_type == 'question' else '(a)', 'printedMarks': marks,
            'choices': [], 'givenData': [], 'contextOrigins': [], 'proofs': [anchor], 'requirements': [requirement]}


def page(items=(), *, start=False, title='Final', role='questions', disposition='readable'):
    return {'disposition': disposition, 'pageRole': role,
            'paper': {'startsPaper': start, 'title': title, 'examType': 'final', 'totalMarks': None,
                      'evidence': 'Original exam header and numbering' if start else 'Continues same header',
                      'confidence': .95, 'uncertain': False},
            'items': list(items), 'sharedContexts': [], 'topics': []}


class Provider:
    def __init__(self, rows, reconciliation=None):
        self.rows, self.reconciliation = rows, reconciliation or rows
        self.calls = []
        self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))

    def create(self, **kwargs):
        self.calls.append(copy.deepcopy(kwargs))
        parts = kwargs['messages'][1]['content']
        request = json.loads(parts[0]['text'])
        assert sum(len(m['content']) if isinstance(m['content'], str) else
                   sum(len(p.get('text', '')) for p in m['content']) for m in kwargs['messages']) <= 24000
        assert sum(p['type'] == 'image_url' for p in parts) <= 8
        rows = self.reconciliation if request['stage'] == 'reconcile' else self.rows
        key = request['source']['ref']['unitId']
        value = rows(request) if callable(rows) else rows[key]
        return SimpleNamespace(choices=[SimpleNamespace(finish_reason='stop', message=SimpleNamespace(content=json.dumps(value)))])


def run(raw, *, kind='historical_exam', file_type='pdf', rows=None, reconciliation=None, **kwargs):
    identity = {**IDENTITY, 'kind': kind}
    manifest = extractor().prepare_pages(identity, raw, file_type)
    rows = rows(manifest) if callable(rows) else rows
    provider = Provider(rows, reconciliation(manifest) if callable(reconciliation) else reconciliation)
    checkpoints = []
    render = (lambda unit: [image_to_image_uri(raw)]) if file_type == 'image' else lambda unit: render_source_unit(raw, file_type, unit['unitId'])
    result = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=reader(), render=render,
                                         checkpoint=checkpoints.append, **kwargs)
    return result, provider, manifest, checkpoints


def repeated_answer_key_text_not_counted_as_question(result):
    # Fixture pages 1..10 each hold one original question; page 11 repeats paper 10's answer.
    origins = [i['supportingSourceRefs'][0]['unitId'] for i in result['items'] if i['itemType'] == 'question']
    return sorted(origins) == [f'p{n:04d}' for n in range(1, 11)] and 'p0011' not in origins


def test_ten_papers_mixed_numbering_and_answer_key_have_distinct_original_identities():
    raw = pdf([f'Final {n}\nQuestion {"I" if n % 2 else "1"}: Explain trees' for n in range(1, 11)] + ['Answers to Question 1: Explain trees'])
    def rows(m):
        result = {}
        for n in range(1, 12):
            questions = [item(proof(m, n))] if n <= 10 else []
            for question in questions:
                question['printedLabel'] = 'I' if n % 2 else '1'
            result[f'p{n:04d}'] = page(questions, start=n <= 10, title=f'Final {min(n,10)}',
                                     role='questions' if n <= 10 else 'answer_key')
        return result
    result, provider, manifest, _ = run(raw, rows=rows)
    assert result['extractionCoverage']['complete'] is True
    assert len(result['papers']) == 10
    assert len({i['id'] for i in result['items']}) == len(result['items']) == 10
    assert repeated_answer_key_text_not_counted_as_question(result)
    assert [i['printedLabel'] for i in result['items']] == ['I', '1']*5
    assert all(i['printedMarks'] is None for i in result['items'])
    assert manifest == validate_supporting_manifest(manifest)
    for stage in ('extract', 'reconcile'):
        requests = [json.loads(c['messages'][1]['content'][0]['text']) for c in provider.calls]
        assert [r['source']['ref']['unitId'] for r in requests if r['stage'] == stage] == [f'p{n:04d}' for n in range(1, 12)]
    next_call = provider.calls[1]['messages'][1]['content']
    assert json.loads(next_call[0]['text'])['previous']['ref']['unitId'] == 'p0001'
    assert next_call[1]['image_url']['url'] == provider.calls[0]['messages'][1]['content'][1]['image_url']['url']


def test_photo_and_vector_pages_keep_original_visual_proofs():
    for raw, file_type in [(png(text='Final exam\nQuestion 1: Explain trees'), 'image'), (pdf(['Question 1: Explain the tree diagram'], vector=True), 'pdf')]:
        result, _, manifest, _ = run(raw, file_type=file_type,
            rows=lambda m: {m['units'][0]['unitId']: page([item(proof(m, visual=True))], start=True)})
        assert result['extractionCoverage']['complete']
        assert result['items'][0]['supportingSourceProofs'][0]['type'] == 'visual_region'
        assert manifest['units'][0]['visualText'] == ''


@pytest.mark.parametrize('file_type', ['docx', 'pptx'])
def test_office_tables_and_all_embedded_assets_are_preserved(file_type):
    out = io.BytesIO()
    if file_type == 'docx':
        doc = Document()
        doc.add_table(rows=1, cols=2).rows[0].cells[0].text = 'Exercise 9: Fill table'
        doc.add_picture(io.BytesIO(png()))
        doc.save(out)
    else:
        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[6])
        slide.shapes.add_table(1, 2, 0, 0, Inches(5), Inches(1)).table.cell(0, 0).text = 'Exercise 9: Fill table'
        for _ in range(10):
            slide.shapes.add_picture(io.BytesIO(png()), 0, Inches(2))
        doc.save(out)
    manifest = extractor().prepare_pages({**IDENTITY, 'kind': 'tutorial'}, out.getvalue(), file_type)
    assert 'Exercise 9: Fill table' in '\n'.join(u['text'] for u in manifest['units'])
    assert sum(len(u['visualAssets']) for u in manifest['units']) == (1 if file_type == 'docx' else 10)
    for unit in manifest['units']:
        assert len(render_source_unit(out.getvalue(), file_type, unit['unitId'])) == len(unit['visualAssets'])


def test_scanned_tutorial_reaches_unique_late_exercise_and_cross_page_subpart():
    raw = pdf(['Course administration: exam-room instructions']*10 + ['Question 1: Unique late problem [10 marks]', '(a) Late subpart'], scanned=True)
    def rows(m):
        anchor = proof(m, 11, visual=True)
        result = {f'p{n:04d}': page([], role='other', disposition='nonacademic') for n in range(1, 11)}
        result['p0011'] = page([item(anchor, 'Unique late problem', marks=10)])
        result['p0012'] = page([item(proof(m, 12, visual=True), 'Late subpart', item_type='subpart',
                                          parent={'origin': anchor, 'itemType': 'question'})])
        return result
    result, _, _, _ = run(raw, kind='tutorial', rows=rows)
    assert result['extractionCoverage']['complete']
    assert len(result['items']) == 2
    late_subpart_id = next(i['id'] for i in result['items'] if i['text'] == 'Late subpart')
    assert late_subpart_id in {i['id'] for i in result['items']}
    child = next(i for i in result['items'] if i['id'] == late_subpart_id)
    assert child['parentId'] == result['items'][0]['id']
    assert child['printedMarks'] is None


@pytest.mark.parametrize('failure', ['unreadable', 'missing_subpart', 'disagreement'])
def test_independent_reconciliation_detects_incomplete_inventory(failure):
    raw = pdf(['Question 1: Explain trees\n(a) Draw tree'])
    def extracted(m):
        return {'p0001': page([item(proof(m))], start=True,
                              disposition='uncertain' if failure == 'unreadable' else 'readable')}
    def reconciled(m):
        rows = extracted(m)
        if failure == 'missing_subpart':
            rows['p0001']['items'].append(item(proof(m, start=26), 'Draw tree', item_type='subpart', parent={'origin': proof(m), 'itemType': 'question'}))
        if failure == 'disagreement':
            rows['p0001']['items'][0]['text'] = 'Explain graphs'
        return rows
    result, _, _, _ = run(raw, rows=extracted, reconciliation=reconciled)
    assert result['extractionCoverage']['complete'] is False
    assert result['extractionCoverage']['unresolvedUnitIds'] == ['p0001']


def test_budget_exhaustion_and_resume_keep_stable_ids_and_spent_calls():
    raw = pdf(['Question 1', 'Question 2'])
    rows = lambda m: {f'p{n:04d}': page([item(proof(m, n), f'Question {n}')], start=n == 1) for n in (1, 2)}
    partial, _, manifest, checkpoints = run(raw, rows=rows, model_call_budget=2)
    assert not partial['extractionCoverage']['complete']
    assert checkpoints[-1]['modelCalls'] == 2
    provider = Provider(rows(manifest))
    done = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda unit: render_source_unit(raw, 'pdf', unit['unitId']), checkpoint=lambda state: None,
        resume=checkpoints[-1], model_call_budget=4, model_calls=2)
    assert done['extractionCoverage']['complete']
    assert [i['id'] for i in done['items']] == [i['id'] for i in partial['items']]
    assert len(provider.calls) == 2


def test_invalid_original_proof_never_becomes_complete():
    raw = pdf(['Question 1'])
    def rows(m):
        bad = proof(m)
        bad['quote'] = 'X'*len(bad['quote'])
        return {'p0001': page([item(bad)], start=True)}
    result, _, _, checkpoints = run(raw, rows=rows)
    assert not result['extractionCoverage']['complete']
    assert checkpoints[-1]['modelCalls'] > 0


def test_long_text_question_keeps_every_character_across_windows():
    doc, out = Document(), io.BytesIO()
    original_text = 'Question 1: ' + 'long original requirement ' * 1400 + 'UNIQUE_FINAL_WORD'
    doc.add_paragraph(original_text)
    doc.save(out)
    def rows(m):
        first = proof(m, end=1)
        def response(request):
            part = request['source']['ranges'][0]
            current = proof(m, start=part['start'], end=part['end'])
            row = item(first, part['text'])
            row['proofs'] = [current]
            row['requirements'] = [{'origin': current, 'text': part['text'], 'category': 'concept', 'proofs': [current]}]
            row['textJoin'] = ''
            return page([row])
        return response
    result, provider, manifest, _ = run(out.getvalue(), kind='tutorial', file_type='docx', rows=rows)
    assert result['extractionCoverage']['complete']
    assert len(result['items']) == 1
    assert result['items'][0]['text'] == original_text
    for stage in ('extract', 'reconcile'):
        ranges = [r['source']['ranges'][0] for call in provider.calls
                  if (r := json.loads(call['messages'][1]['content'][0]['text']))['stage'] == stage]
        assert ''.join(r['text'] for r in ranges) == manifest['units'][0]['text']
        assert ranges[0]['start'] == 0 and ranges[-1]['end'] == len(original_text)


def test_standalone_heading_has_own_requirement_beside_unrelated_question():
    raw = pdf(['Graph theory\nQuestion 1: Explain trees'])
    def rows(m):
        heading = item(proof(m, end=12), 'Graph theory', item_type='topic_statement')
        question = item(proof(m, start=13), 'Explain trees')
        response = page([heading, question], start=True)
        response['topics'] = [{'text': 'Graph theory', 'items': [{'origin': heading['origin'], 'itemType': 'topic_statement'}],
            'requirements': [{k: heading['requirements'][0][k] for k in ('origin', 'text', 'category')}], 'proofs': heading['proofs']}]
        return {'p0001': response}
    result, _, _, _ = run(raw, rows=rows)
    assert result['extractionCoverage']['complete']
    topic, heading, question = result['topics'][0], result['items'][0], result['items'][1]
    assert topic['requirementIds'] == [heading['requirements'][0]['id']]
    assert question['id'] not in topic['itemIds']
    assert result['papers'][0]['totalQuestions'] == 1


def test_shared_context_and_continued_parent_keep_original_identity():
    raw = pdf(['Use table A\nQuestion 1: Given data', 'continue question\n(a) Calculate'])
    def rows(m):
        context = {'origin': proof(m, end=11), 'text': 'Use table A', 'proofs': [proof(m, end=11)]}
        first = item(proof(m, start=12), 'Given data', marks=10)
        first['contextOrigins'] = [context['origin']]
        a = page([first], start=True)
        a['sharedContexts'] = [context]
        continued = item(first['origin'], 'continue question', marks=10)
        continued['proofs'] = [proof(m, 2, end=17)]
        continued['requirements'] = [{'origin': proof(m, 2, end=17), 'text': 'Continue operation', 'category': 'operation', 'proofs': continued['proofs']}]
        continued['contextOrigins'] = [context['origin']]
        subpart = item(proof(m, 2, start=18), 'Calculate', item_type='subpart', parent={'origin': first['origin'], 'itemType': 'question'})
        return {'p0001': a, 'p0002': page([continued, subpart])}
    result, _, _, _ = run(raw, rows=rows)
    assert result['extractionCoverage']['complete']
    assert len(result['items']) == 2
    assert result['items'][0]['text'] == 'Given data\ncontinue question'
    assert result['items'][0]['sharedContextIds'] == [result['sharedContexts'][0]['id']]
    assert result['items'][1]['printedMarks'] is None


def test_returned_success_is_checkpointed_before_parsing_and_recovered():
    raw = pdf(['Question 1'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    rows = {'p0001': page([item(proof(manifest))], start=True)}
    provider = Provider(rows)
    checkpoints = []
    def interrupt(state):
        checkpoints.append(state)
        if state['pending'] is not None:
            raise InterruptedError('Process interrupted after provider success')
    with pytest.raises(InterruptedError):
        extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
            render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=interrupt)
    assert checkpoints[-1]['modelCalls'] == 1
    assert checkpoints[-1]['cursor'] == 0
    recovered = Provider(rows)
    result = extractor(recovered).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None, resume=checkpoints[-1])
    assert result['extractionCoverage']['complete']
    assert len(recovered.calls) == 1  # independent reconciliation, pending extraction not sent again
    assert result['extractionCoverage']['modelCalls'] == 2


def test_truncated_and_invalid_completions_spend_budget():
    raw = pdf(['Question 1'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    class Broken:
        def __init__(self):
            self.chat = SimpleNamespace(completions=SimpleNamespace(create=self.create))
            self.count = 0
        def create(self, **kwargs):
            self.count += 1
            return SimpleNamespace(choices=[SimpleNamespace(finish_reason='length' if self.count == 1 else 'stop',
                message=SimpleNamespace(content='{broken'))])
    checkpoints, provider = [], Broken()
    result = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=checkpoints.append, model_call_budget=2)
    assert not result['extractionCoverage']['complete']
    assert checkpoints[-1]['modelCalls'] == provider.count == 2
    assert [c['modelCalls'] for c in checkpoints if c['pending'] is not None] == [1, 2]


def test_checkpoint_rejects_changed_source_skipped_window_and_forged_proof():
    raw = pdf(['Question 1', 'Question 2'])
    rows = lambda m: {f'p{n:04d}': page([item(proof(m, n))], start=n == 1) for n in (1, 2)}
    _, _, manifest, checkpoints = run(raw, rows=rows, model_call_budget=1)
    for mutation in ('source', 'cursor', 'proof'):
        state = copy.deepcopy(checkpoints[-1])
        if mutation == 'source':
            state['sourceRevision'] = 'b'*64
        elif mutation == 'cursor':
            state['cursor'] = 2
        else:
            state['responses']['extract']['0']['items'][0]['proofs'][0]['quote'] = 'forged'
        with pytest.raises(ValueError):
            extractor(Provider(rows(manifest))).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
                render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None, resume=state)


def test_multiframe_image_is_explicitly_rejected_instead_of_reading_only_first_frame():
    out = io.BytesIO()
    Image.new('RGB', (30, 30), 'white').save(out, format='GIF', save_all=True,
                                          append_images=[Image.new('RGB', (30, 30), 'black')])
    with pytest.raises(ValueError, match='(?i)multiple|frames|PDF'):
        extractor().prepare_pages(IDENTITY, out.getvalue(), 'image')


def test_unsupported_office_visual_requires_conversion():
    from pptx.enum.shapes import MSO_SHAPE
    doc, out = Presentation(), io.BytesIO()
    slide = doc.slides.add_slide(doc.slide_layouts[6])
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
    doc.save(out)
    with pytest.raises(ValueError, match='(?i)PDF'):
        extractor().prepare_pages({**IDENTITY, 'kind': 'tutorial'}, out.getvalue(), 'pptx')


def test_missing_render_asset_is_incomplete_without_provider_calls():
    raw = pdf(['Question 1'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    provider = Provider({})
    result = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: [], checkpoint=lambda s: None)
    assert not result['extractionCoverage']['complete']
    assert not provider.calls


def test_observed_question_formats_are_per_paper_and_unknown_stays_unknown():
    raw = pdf(['Quiz: Choose A or B', 'Final: prove theorem', 'Question of unclear format'])
    def rows(m):
        result = {}
        for n, question_type in enumerate(('mcq', 'proof', None), 1):
            question = item(proof(m, n))
            question['questionType'] = question_type
            row = page([question], start=n < 3, title='Quiz' if n == 1 else 'Final')
            row['paper']['examType'] = 'quiz' if n == 1 else 'final'
            result[f'p{n:04d}'] = row
        return result
    result, _, _, _ = run(raw, rows=rows)
    assert result['extractionCoverage']['complete']
    assert [i['questionType'] for i in result['items']] == ['mcq', 'proof', None]
    assert result['papers'][0]['questionTypes'] == [{'type': 'mcq', 'count': 1, 'questionIds': [result['items'][0]['id']]}]
    assert result['papers'][1]['questionTypes'] == [{'type': 'proof', 'count': 1, 'questionIds': [result['items'][1]['id']]}]
    assert result['papers'][1]['unknownQuestionTypeIds'] == [result['items'][2]['id']]
    assert [p['examType'] for p in result['papers']] == ['quiz', 'final']
    assert all(p['supportingSourceRefs'] for p in result['papers'])


def test_more_than_eight_office_assets_are_all_submitted_and_reconciled():
    doc, out = Presentation(), io.BytesIO()
    slide = doc.slides.add_slide(doc.slide_layouts[6])
    for n in range(12):
        slide.shapes.add_picture(io.BytesIO(png((n*20, 40, 90))), 0, Inches(2))
    doc.save(out)
    def rows(m):
        def response(request):
            questions = []
            for asset in request['source']['assetIds']:
                p = proof(m, visual=True)
                p['assetId'] = asset
                questions.append(item(p, 'Original exercise ' + asset))
            return page(questions)
        return response
    result, provider, manifest, _ = run(out.getvalue(), kind='tutorial', file_type='pptx', rows=rows)
    assert result['extractionCoverage']['complete']
    assert len(result['items']) == 12
    expected = [a['assetId'] for a in manifest['units'][0]['visualAssets']]
    for stage in ('extract', 'reconcile'):
        assert [asset for c in provider.calls if (r := json.loads(c['messages'][1]['content'][0]['text']))['stage'] == stage
                for asset in r['source']['assetIds']] == expected


def test_original_dispositions_and_answer_key_role_survive_in_coverage():
    raw = pdf(['Question 1', 'Answers 1'])
    result, _, _, _ = run(raw, rows=lambda m: {
        'p0001': page([item(proof(m))], start=True), 'p0002': page([], role='answer_key')})
    assert result['extractionCoverage']['complete']
    assert result['extractionCoverage']['pageRoles']['p0002'] == ['answer_key']
    assert result['extractionCoverage']['dispositions']['p0001'] == ['readable']


def test_invalid_low_spent_count_resume_is_rejected_without_explicit_fresh_budget():
    raw = pdf(['Question 1'])
    _, _, manifest, checkpoints = run(raw, rows=lambda m: {'p0001': page([item(proof(m))], start=True)}, model_call_budget=1)
    state = copy.deepcopy(checkpoints[-1])
    state['modelCalls'] = 0
    with pytest.raises(ValueError, match='(?i)call|checkpoint'):
        extractor(Provider({})).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
            render=lambda u: [], checkpoint=lambda s: None, resume=state)


def test_whitespace_only_readable_response_cannot_hide_academic_page():
    raw = pdf(['Question 1'])
    result, _, _, _ = run(raw, rows={'p0001': page([], start=True, role='other')})
    assert not result['extractionCoverage']['complete']


def test_duplicate_origin_in_one_response_cannot_hide_missing_question():
    raw = pdf(['Question 1'])
    def rows(m):
        q = item(proof(m))
        return {'p0001': page([q, copy.deepcopy(q)], start=True)}
    result, _, _, _ = run(raw, rows=rows)
    assert not result['extractionCoverage']['complete']


def test_report_coverage_ledgers_store_ranges_not_duplicate_original_text():
    raw = pdf(['Question 1: SOURCE_TEXT_MUST_NOT_BE_COPIED_IN_COVERAGE'])
    result, _, _, _ = run(raw, rows=lambda m: {'p0001': page([item(proof(m))], start=True)})
    submitted = result['extractionCoverage']['submitted']
    assert submitted['extract'][0]['ranges'] == [{'field': 'text', 'start': 0, 'end': 54}]
    assert 'SOURCE_TEXT_MUST_NOT_BE_COPIED_IN_COVERAGE' not in json.dumps(submitted)


def test_explicit_retry_rechecks_disagreement_without_repeating_valid_prefix():
    raw = pdf(['Question 1', 'Question 2'])
    def rows(m):
        return {f'p{n:04d}': page([item(proof(m, n), f'Question {n}')], start=n == 1) for n in (1, 2)}
    def wrong(m):
        values = rows(m)
        values['p0002']['items'][0]['text'] = 'Uncertain wording'
        return values
    result, _, manifest, states = run(raw, rows=rows, reconciliation=wrong)
    assert not result['extractionCoverage']['complete']
    provider = Provider(rows(manifest))
    retried = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None,
        resume=states[-1], model_calls=0)
    assert retried['extractionCoverage']['complete']
    extract_requests = [json.loads(c['messages'][1]['content'][0]['text']) for c in provider.calls]
    assert [r['source']['ref']['unitId'] for r in extract_requests if r['stage'] == 'extract'] == ['p0002']


def test_duplicate_atomic_requirement_cannot_produce_complete_unpublishable_report():
    raw = pdf(['Question 1'])
    def rows(m):
        q = item(proof(m))
        q['requirements'].append(copy.deepcopy(q['requirements'][0]))
        return {'p0001': page([q], start=True)}
    result, _, _, _ = run(raw, rows=rows)
    assert not result['extractionCoverage']['complete']


def test_single_question_source_proof_outside_supplied_range_is_rejected():
    out, doc = io.BytesIO(), Document()
    doc.add_paragraph('Start' + 'z'*10000 + 'End')
    doc.save(out)
    result, _, _, _ = run(out.getvalue(), kind='tutorial', file_type='docx',
        rows=lambda m: {'b0001': page([item(proof(m))])})
    assert not result['extractionCoverage']['complete']


def test_ambiguous_first_boundary_does_not_invent_one_confident_paper():
    raw = pdf(['Continued Question 2'])
    result, _, _, _ = run(raw, rows=lambda m: {'p0001': page([item(proof(m))])})
    assert not result['extractionCoverage']['complete']
    assert not result['papers']


def test_duplicate_requirements_across_parent_child_remain_incomplete():
    raw = pdf(['Question 1 (a) Define tree'])
    def rows(m):
        parent = item(proof(m))
        child = item(proof(m), item_type='subpart', parent={'origin': proof(m), 'itemType': 'question'})
        return {'p0001': page([parent, child], start=True)}
    result, _, _, _ = run(raw, rows=rows)
    assert not result['extractionCoverage']['complete']


def test_explicit_fresh_budget_retains_valid_extraction_when_only_reconciliation_is_pending():
    raw = pdf(['Question 1', 'Question 2'])
    rows = lambda m: {f'p{n:04d}': page([item(proof(m, n))], start=n == 1) for n in (1, 2)}
    _, _, manifest, states = run(raw, rows=rows, model_call_budget=2)
    provider = Provider(rows(manifest))
    result = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None,
        resume=states[-1], model_calls=0, model_call_budget=2)
    assert result['extractionCoverage']['complete']
    assert all(json.loads(c['messages'][1]['content'][0]['text'])['stage'] == 'reconcile' for c in provider.calls)


def test_long_pdf_text_and_boundary_context_are_fully_read_without_old_page_limit():
    doc = fitz.open()
    page_ = doc.new_page(width=600, height=6000)
    original = 'Question 1\n' + '\n'.join(f'Unique instruction line {n:04d}: compare trees and graphs.' for n in range(320))
    page_.insert_text((40, 50), original, fontsize=10)
    raw = doc.tobytes()
    doc.close()
    def rows(m):
        first = proof(m, end=1)
        def response(request):
            part = request['source']['ranges'][0]
            current = proof(m, start=part['start'], end=part['end'])
            q = item(first, part['text'])
            q['textJoin'] = ''
            q['proofs'] = [current]
            q['requirements'] = [{'origin': current, 'text': part['text'], 'category': 'operation', 'proofs': [current]}]
            return page([q], start=part['start'] == 0)
        return response
    result, _, manifest, _ = run(raw, rows=rows)
    assert len(manifest['units'][0]['text']) > 12000
    assert result['extractionCoverage']['complete']
    assert result['items'][0]['text'] == manifest['units'][0]['text']
    assert len(result['papers']) == 1


@pytest.mark.parametrize('bad', [None, 4, [], {'items': [None]}])
def test_malformed_response_shapes_spend_bounded_calls_and_stay_incomplete(bad):
    raw = pdf(['Question 1'])
    result, _, _, states = run(raw, rows={'p0001': bad})
    assert not result['extractionCoverage']['complete']
    assert states[-1]['modelCalls'] <= 2


def test_provider_transport_failure_preserves_incomplete_checkpoint():
    raw = pdf(['Question 1'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    class TransportFailure(Exception):
        status_code = 400
    def fail(**kwargs):
        raise TransportFailure('Images unsupported by provider')
    client = SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=fail)))
    states = []
    result = extractor(client).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=states.append)
    assert not result['extractionCoverage']['complete']
    assert states[-1]['modelCalls'] == 0
    assert 'provider' in states[-1]['error'].lower()


def test_default_source_reader_uses_available_request_budget_without_6000_character_cap():
    from src.services.source_evidence import SourceEvidence
    from .test_source_context import source_scope
    raw = pdf(['Question 1: Explain trees'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    scope = source_scope(['Algorithms ' + 'original lecture explanation '*400])
    scope['documents'][0]['guide']['sections'][0]['purpose'] = 'Detailed navigation hint '*280
    original_scope = copy.deepcopy(scope)
    source_reader = SourceEvidence(scope, object(), 'u', 'c', None, 'no-lecture-model-calls')
    provider = Provider({'p0001': page([item(proof(manifest))], start=True)})
    result = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=source_reader,
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None)
    assert result['extractionCoverage']['complete']
    packets = [json.loads(c['messages'][1]['content'][0]['text'])['lectureNavigation'] for c in provider.calls]
    assert all(len(json.dumps(p, ensure_ascii=False, separators=(',', ':'))) > 6000 for p in packets)
    assert packets[0]['guide'][0]['guide']['sections'][0]['purpose'] == 'Detailed navigation hint '*280
    assert source_reader.scope == original_scope
    assert source_reader.text_budget == 30000 and source_reader.image_budget == 8
    assert source_reader.supplied == [] and source_reader.completed_batches == 0


def test_default_reader_no_course_title_match_still_supplies_original_partial_navigation():
    from src.services.source_evidence import SourceEvidence
    from .test_source_context import source_scope
    raw = pdf(['Question 1'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    scope = source_scope(['Tree definition and traversal'])
    source_reader = SourceEvidence(scope, object(), 'u', 'c', None, 'no-lecture-model-calls')
    provider = Provider({'p0001': page([item(proof(manifest))], start=True)})
    result = extractor(provider).extract(manifest, course_title='CS101', lecture_reader=source_reader,
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None)
    assert result['extractionCoverage']['complete']
    packet = json.loads(provider.calls[0]['messages'][1]['content'][0]['text'])['lectureNavigation']
    assert packet['coverage']['complete'] is False and packet['coverage']['mode'] == 'partial'
    assert 'Tree definition' in json.dumps(packet['evidence'])


def test_navigation_selects_guide_rows_for_supplied_units_without_changing_full_scope():
    from src.services.source_evidence import SourceEvidence
    from .test_source_context import source_scope
    raw = pdf(['Question 1'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    scope = source_scope(['Algorithms about trees', 'Unrelated guide example'])
    scope['documents'][0]['guide']['sections'].append({'label': 'Other unit', 'purpose': 'unrelated '*4000, 'unitIds': ['p0002']})
    source_reader = SourceEvidence(scope, object(), 'u', 'c', None, 'no-lecture-model-calls')
    provider = Provider({'p0001': page([item(proof(manifest))], start=True)})
    result = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=source_reader,
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None)
    assert result['extractionCoverage']['complete']
    packet = json.loads(provider.calls[0]['messages'][1]['content'][0]['text'])['lectureNavigation']
    assert [r['label'] for r in packet['guide'][0]['guide']['sections']] == ['Only guide heading']
    assert len(source_reader.scope['documents'][0]['guide']['sections']) == 2
    assert packet['navigationView'] == 'partial'


def test_structured_retry_reuses_full_available_navigation_budget():
    from src.services.source_evidence import SourceEvidence
    from .test_source_context import source_scope
    raw = pdf(['Question 1'])
    manifest = extractor().prepare_pages(IDENTITY, raw, 'pdf')
    scope = source_scope(['Algorithms definition'])
    scope['documents'][0]['guide']['sections'][0]['purpose'] = 'Original navigation '*600
    source_reader = SourceEvidence(scope, object(), 'u', 'c', None, 'no-lecture-model-calls')
    provider = Provider({'p0001': page([item(proof(manifest))], start=True)})
    calls = 0
    original_create = provider.create
    def create(**kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            return SimpleNamespace(choices=[SimpleNamespace(finish_reason='length', message=SimpleNamespace(content='{'))])
        return original_create(**kwargs)
    provider.chat.completions.create = create
    result = extractor(provider).extract(manifest, course_title='Algorithms', lecture_reader=source_reader,
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None)
    assert result['extractionCoverage']['complete']
    assert result['extractionCoverage']['modelCalls'] == 3


@pytest.mark.parametrize('recover', [False, True])
def test_same_atomic_requirement_continued_across_pages_unions_evidence_once(recover):
    from src.supporting.contracts import validate_alignment_report
    raw = pdf(['Question 1: Explain trees', 'Continued explanation of trees'])
    def rows(m):
        first = item(proof(m), 'Explain trees')
        continued = item(first['origin'], 'Continued explanation')
        continued['proofs'] = [proof(m, 2)]
        continued['requirements'] = copy.deepcopy(first['requirements'])
        continued['requirements'][0]['proofs'] = [proof(m, 2)]
        return {'p0001': page([first], start=True), 'p0002': page([continued])}
    result, _, manifest, states = run(raw, rows=rows, model_call_budget=2 if recover else 512)
    if recover:
        result = extractor(Provider(rows(manifest))).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
            render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None,
            resume=states[-1], model_call_budget=4)
    assert result['extractionCoverage']['complete']
    requirements = result['items'][0]['requirements']
    assert len(requirements) == len({r['id'] for r in requirements}) == 1
    assert [r['unitId'] for r in requirements[0]['supportingSourceRefs']] == ['p0001', 'p0002']
    assert requirements[0]['supportingSourceProofs'] == [proof(manifest), proof(manifest, 2)]
    # The real report validator must accept the extracted inventory before Task5
    # has performed alignment; empty alignment coverage is explicitly incomplete.
    report = {**result, 'version': 1, 'alignmentVersion': 1, 'sourceRevision': manifest['sourceRevision'],
        'scopeSnapshot': reader().snapshot(), 'alignments': [], 'alignmentCoverage': {
            'itemIds': [i['id'] for i in result['items']], 'checkedItemIds': [],
            'unresolvedItemIds': [i['id'] for i in result['items']], 'complete': False}}
    assert validate_alignment_report(report)['items'][0]['requirements'] == requirements


def test_continued_requirement_disagreement_is_not_hidden_by_evidence_union():
    raw = pdf(['Question 1: Explain trees', 'Continued explanation of trees'])
    def rows(m):
        q = item(proof(m), 'Explain trees')
        continued = copy.deepcopy(q)
        continued['proofs'] = [proof(m, 2)]
        continued['requirements'][0]['proofs'] = [proof(m, 2)]
        return {'p0001': page([q], start=True), 'p0002': page([continued])}
    def disagreement(m):
        result = rows(m)
        result['p0002']['items'][0]['requirements'][0]['text'] = 'Explain graphs'
        return result
    result, _, _, _ = run(raw, rows=rows, reconciliation=disagreement)
    assert not result['extractionCoverage']['complete']
    assert 'p0002' in result['extractionCoverage']['unresolvedUnitIds']


@pytest.mark.parametrize('target,field', [('item', 'printedMarks'), ('item', 'printedLabel'), ('item', 'parent'),
    ('item', 'questionType'), ('paper', 'totalMarks'), ('paper', 'examType')])
def test_missing_nullable_fields_normalize_before_accepted_checkpoint(target, field):
    raw = pdf(['Question 1'])
    def rows(m):
        value = page([item(proof(m))], start=True)
        record = value['items'][0] if target == 'item' else value['paper']
        record.pop(field, None)
        return {'p0001': value}
    result, _, manifest, states = run(raw, rows=rows)
    assert result['extractionCoverage']['complete']
    for state in states:
        for response in state['responses']['extract'].values():
            row = response['items'][0] if target == 'item' else response['paper']
            assert field in row and row[field] is None
    recovered = extractor(Provider({})).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None, resume=states[-1])
    assert recovered['extractionCoverage']['complete']


@pytest.mark.parametrize('malformation', ['topic_item_null', 'topic_item_list', 'topic_requirement_null',
    'topic_requirement_list', 'topic_origin_null', 'topic_requirement_origin_null', 'item_origin_null',
    'requirement_null', 'requirement_origin_list', 'parent_origin_null', 'context_origin_null', 'proof_ref_null',
    'topic_item_missing_type', 'topic_requirement_missing_text', 'topic_requirement_missing_category',
    'topic_requirement_missing_origin', 'context_missing_text', 'item_missing_choices',
    'paper_missing_confidence', 'requirement_missing_proofs', 'proof_missing_box'])
def test_malformed_nested_inventory_is_counted_rejected_and_recoverable(malformation):
    raw = pdf(['Question 1: Explain trees'])
    def valid(m):
        question = item(proof(m))
        value = page([question], start=True)
        value['topics'] = [{'text': 'Explain trees', 'items': [{'origin': proof(m), 'itemType': 'question'}],
            'requirements': [{k: question['requirements'][0][k] for k in ('origin', 'text', 'category')}], 'proofs': [proof(m)]}]
        return {'p0001': value}
    def broken(m):
        values = valid(m)
        value = values['p0001']
        topic, q = value['topics'][0], value['items'][0]
        if malformation.startswith('topic_item_') and 'missing' not in malformation:
            topic['items'] = [None if malformation.endswith('null') else []]
        elif malformation.startswith('topic_requirement_') and 'origin' not in malformation and 'missing' not in malformation:
            topic['requirements'] = [None if malformation.endswith('null') else []]
        elif malformation == 'topic_origin_null': topic['items'][0]['origin'] = None
        elif malformation == 'topic_requirement_origin_null': topic['requirements'][0]['origin'] = None
        elif malformation == 'item_origin_null': q['origin'] = None
        elif malformation == 'requirement_null': q['requirements'] = [None]
        elif malformation == 'requirement_origin_list': q['requirements'][0]['origin'] = []
        elif malformation == 'parent_origin_null': q['parent'] = {'itemType': 'question', 'origin': None}
        elif malformation == 'context_origin_null': q['contextOrigins'] = [None]
        elif malformation == 'proof_ref_null': q['proofs'] = [{**proof(m), 'ref': None}]
        elif malformation == 'topic_item_missing_type': topic['items'][0].pop('itemType')
        elif malformation.startswith('topic_requirement_missing_'): topic['requirements'][0].pop(malformation.rsplit('_', 1)[-1])
        elif malformation == 'context_missing_text': value['sharedContexts'] = [{'origin': proof(m), 'proofs': [proof(m)]}]
        elif malformation == 'item_missing_choices': q.pop('choices')
        elif malformation == 'paper_missing_confidence': value['paper'].pop('confidence')
        elif malformation == 'requirement_missing_proofs': q['requirements'][0].pop('proofs')
        elif malformation == 'proof_missing_box':
            p = proof(m, visual=True)
            p.pop('box')
            q['proofs'] = [p]
        return values
    result, _, manifest, states = run(raw, rows=broken, model_call_budget=2)
    assert not result['extractionCoverage']['complete']
    assert states[-1]['modelCalls'] == 2
    assert states[-1]['responses'] == {'extract': {}, 'reconcile': {}}
    recovered = extractor(Provider(valid(manifest))).extract(manifest, course_title='Algorithms', lecture_reader=reader(),
        render=lambda u: render_source_unit(raw, 'pdf', u['unitId']), checkpoint=lambda s: None,
        resume=states[-1], model_calls=0, model_call_budget=2)
    assert recovered['extractionCoverage']['complete']
